#!/usr/bin/env python3
"""Convert markdown documentation to DOCX format."""

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
import re


def parse_markdown_to_docx(md_content, output_path):
    """Convert markdown content to DOCX document."""
    doc = Document()
    
    lines = md_content.split('\n')
    i = 0
    
    while i < len(lines):
        line = lines[i]
        
        # Skip empty markdown markers
        if line.strip().startswith('```'):
            i += 1
            continue
            
        # Handle headings
        if line.startswith('# '):
            heading = line.replace('# ', '').strip()
            p = doc.add_heading(heading, level=1)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            i += 1
        elif line.startswith('## '):
            heading = line.replace('## ', '').strip()
            p = doc.add_heading(heading, level=2)
            i += 1
        elif line.startswith('### '):
            heading = line.replace('### ', '').strip()
            p = doc.add_heading(heading, level=3)
            i += 1
        # Handle bullet points
        elif line.startswith('- '):
            text = line.replace('- ', '').strip()
            p = doc.add_paragraph(text, style='List Bullet')
            i += 1
        # Handle code blocks
        elif line.strip().startswith('{') or line.strip().startswith('['):
            # Collect code block lines
            code_lines = [line]
            i += 1
            while i < len(lines) and not lines[i].strip().startswith('```'):
                code_lines.append(lines[i])
                i += 1
            code_text = '\n'.join(code_lines).strip()
            if code_text:
                p = doc.add_paragraph(code_text, style='Normal')
                p_format = p.paragraph_format
                p_format.left_indent = Inches(0.5)
                for run in p.runs:
                    run.font.name = 'Courier New'
                    run.font.size = Pt(9)
            i += 1
        # Handle empty lines
        elif not line.strip():
            i += 1
        # Handle regular paragraphs
        else:
            # Collect paragraph lines
            para_lines = []
            while i < len(lines) and lines[i].strip() and not lines[i].startswith('#') and not lines[i].startswith('- '):
                para_lines.append(lines[i].strip())
                i += 1
            
            if para_lines:
                text = ' '.join(para_lines)
                if text:
                    p = doc.add_paragraph(text)
    
    doc.save(output_path)
    print(f"✓ Created {output_path}")


# Read and convert PRD
with open('/home/mntel/Desktop/projects/culture/PRD.md', 'r') as f:
    prd_content = f.read()

parse_markdown_to_docx(prd_content, '/home/mntel/Desktop/projects/culture/PRD.docx')

# Read and convert FRD
with open('/home/mntel/Desktop/projects/culture/FRD.md', 'r') as f:
    frd_content = f.read()

parse_markdown_to_docx(frd_content, '/home/mntel/Desktop/projects/culture/FRD.docx')

print("\n✓ Conversion complete!")
print("  - PRD.docx")
print("  - FRD.docx")
